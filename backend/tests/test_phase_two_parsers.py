"""Parser contracts for phase-two document structure metadata."""

import os
import tempfile
import unittest

from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app.models  # noqa: F401 - register model metadata
from app.collector.parsers.docx_parser import DocxParser
from app.collector.parsers.pdf_parser import PDFParser
from app.collector.parsers.pptx_parser import PptxParser
from app.collector.pipeline import _XlsxParser
from app.models.base import Base
from app.models.chat import DocumentChunk
from app.models.document import Document
from app.models.workspace import Workspace
from app.services.hybrid_retrieval import upsert_document_chunks
from tests.support import create_user, create_workspace


class ParserStructureMetadataTests(unittest.TestCase):
    def test_parsers_preserve_page_numbers_and_heading_paths(self):
        """Missing structural metadata would make the assertions below fail."""
        import fitz
        import openpyxl
        from docx import Document as DocxDocument
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            pdf_path = os.path.join(directory, "book.pdf")
            pdf = fitz.open()
            for text in ("Cover", "Chapter body", "Section body"):
                page = pdf.new_page()
                page.insert_text((72, 72), text)
            pdf.set_toc([[1, "第一章", 2], [2, "1.1 小节", 3]])
            pdf.save(pdf_path)
            pdf.close()

            docx_path = os.path.join(directory, "book.docx")
            docx = DocxDocument()
            docx.add_heading("第一章", level=1)
            docx.add_paragraph("Chapter body")
            docx.add_heading("1.1 小节", level=2)
            docx.add_paragraph("Section body")
            docx.save(docx_path)

            pptx_path = os.path.join(directory, "slides.pptx")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
            title.text = "第一章"
            body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
            body.text = "Slide body"
            presentation.save(pptx_path)

            xlsx_path = os.path.join(directory, "sheet.xlsx")
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "Sheet1"
            worksheet.append(["Name", "Score"])
            workbook.save(xlsx_path)
            workbook.close()

            pdf_chunks = PDFParser().parse(pdf_path)
            docx_chunks = DocxParser().parse(docx_path)
            pptx_chunks = PptxParser().parse(pptx_path)
            xlsx_chunks = _XlsxParser().parse(xlsx_path)

        self.assertEqual(pdf_chunks[0]["metadata"]["section_path"], [])
        self.assertEqual(pdf_chunks[1]["metadata"], {
            "page_num": 2,
            "heading": "第一章",
            "heading_level": 1,
            "section_path": ["第一章"],
            "source_file": "book.pdf",
        })
        self.assertEqual(pdf_chunks[2]["metadata"]["section_path"], ["第一章", "1.1 小节"])
        self.assertIsNone(docx_chunks[1]["metadata"]["page_num"])
        self.assertEqual(docx_chunks[1]["metadata"]["heading_level"], 2)
        self.assertEqual(docx_chunks[1]["metadata"]["section_path"], ["第一章", "1.1 小节"])
        self.assertEqual(pptx_chunks[0]["metadata"]["page_num"], 1)
        self.assertEqual(pptx_chunks[0]["metadata"]["heading_level"], 1)
        self.assertEqual(pptx_chunks[0]["metadata"]["section_path"], ["第一章"])
        self.assertEqual(xlsx_chunks[0]["metadata"]["section_path"], ["Sheet1"])
        self.assertIn("Name | Score", xlsx_chunks[0]["content"])

    def test_empty_title_placeholder_does_not_promote_slide_body_to_heading(self):
        """Treating a content placeholder as a title would assign false hierarchy."""
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            pptx_path = os.path.join(directory, "untitled-content-slide.pptx")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.placeholders[1].text = "Slide body"
            presentation.save(pptx_path)

            chunks = PptxParser().parse(pptx_path)

        self.assertEqual(chunks[0]["metadata"]["heading"], "")
        self.assertIsNone(chunks[0]["metadata"]["heading_level"])
        self.assertEqual(chunks[0]["metadata"]["section_path"], [])

    def test_docx_preserves_heading_levels_through_nine(self):
        """Ignoring Heading 7 or Heading 9 would collapse deep sections into body text."""
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as directory:
            docx_path = os.path.join(directory, "deep-headings.docx")
            document = DocxDocument()
            document.add_heading("Chapter 1", level=1)
            document.add_paragraph("Chapter body")
            document.add_heading("Deep heading", level=7)
            document.add_paragraph("Deep body")
            document.add_heading("Terminal heading", level=9)
            document.add_paragraph("Terminal body")
            document.save(docx_path)

            chunks = DocxParser().parse(docx_path)

        self.assertEqual([chunk["metadata"]["heading_level"] for chunk in chunks], [1, 7, 9])
        self.assertEqual(chunks[1]["metadata"]["section_path"], ["Chapter 1", "Deep heading"])
        self.assertEqual(
            chunks[2]["metadata"]["section_path"],
            ["Chapter 1", "Deep heading", "Terminal heading"],
        )

    def test_docx_replaces_skipped_heading_level_without_retaining_deeper_paths(self):
        """Slicing by list length leaves old deep headings after a skipped-level fallback."""
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as directory:
            docx_path = os.path.join(directory, "skipped-level-fallback.docx")
            document = DocxDocument()
            document.add_heading("Chapter 1", level=1)
            document.add_heading("Deep branch A", level=7)
            document.add_heading("Terminal branch", level=9)
            document.add_paragraph("Terminal body")
            document.add_heading("Deep branch B", level=7)
            document.add_paragraph("Replacement body")
            document.save(docx_path)

            chunks = DocxParser().parse(docx_path)

        self.assertEqual(
            chunks[-1]["metadata"]["section_path"],
            ["Chapter 1", "Deep branch B"],
        )


class ChunkStructurePersistenceTests(unittest.IsolatedAsyncioTestCase):
    async def test_upsert_persists_heading_level_and_json_section_path(self):
        """Dropping hierarchy while indexing would make the stored row lose metadata."""
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        try:
            async with session_factory() as db:
                user = await create_user(db)
                workspace = await create_workspace(
                    db, user, name="Parser Test", slug="parser-test"
                )
                document = Document(
                    workspace_id=workspace.id,
                    filename="book.pdf",
                    file_path="/tmp/book.pdf",
                    file_type="pdf",
                )
                db.add(document)
                await db.flush()

                await upsert_document_chunks(
                    db,
                    workspace.id,
                    document.id,
                    document.filename,
                    [{
                        "content": "Section body",
                        "metadata": {
                            "page_num": 2,
                            "heading": "1.1 小节",
                            "heading_level": 2,
                            "section_path": ["第一章", "1.1 小节"],
                        },
                    }],
                )

                stored = (await db.execute(select(DocumentChunk))).scalar_one()
                self.assertEqual(stored.heading_level, 2)
                self.assertEqual(stored.section_path, ["第一章", "1.1 小节"])
        finally:
            await engine.dispose()

    async def test_pipeline_serializes_section_path_for_chroma_metadata(self):
        """Chroma metadata must not receive a list even when DB metadata retains one."""
        class Parser:
            def parse(self, _file_path):
                return [{
                    "content": "Section body",
                    "metadata": {
                        "page_num": 1,
                        "heading": "第一章",
                        "heading_level": 1,
                        "section_path": ["第一章"],
                    },
                }]

        class Splitter:
            def split_documents(self, chunks):
                return chunks

        class Embeddings:
            async def embed_texts(self, _texts):
                return [[0.0, 1.0]]

        class VectorStore:
            def __init__(self):
                self.metadatas = []

            async def replace_document(self, *, metadatas, **_kwargs):
                self.metadatas = metadatas

        from app.collector.pipeline import DocumentPipeline

        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)
        try:
            vector_store = VectorStore()
            pipeline = DocumentPipeline(Embeddings(), vector_store, Splitter())
            pipeline._get_parser = lambda _file_type: Parser()  # type: ignore[method-assign]
            async with session_factory() as db:
                user = await create_user(db)
                workspace = await create_workspace(
                    db, user, name="Pipeline Test", slug="pipeline-test"
                )
                document = Document(
                    workspace_id=workspace.id,
                    filename="book.pdf",
                    file_path="ignored.pdf",
                    file_type="pdf",
                )
                db.add(document)
                await db.commit()

                await pipeline.process_document(
                    document.id,
                    document.file_path,
                    document.file_type,
                    workspace.id,
                    db,
                )

                rows = (await db.execute(select(DocumentChunk))).scalars().all()
                stored = next(row for row in rows if row.chunk_level == "child")
                parents = [row for row in rows if row.chunk_level == "parent"]
                self.assertEqual(vector_store.metadatas[0]["section_path"], '["第一章"]')
                self.assertEqual(stored.section_path, ["第一章"])
                self.assertEqual(len(parents), 1)
                self.assertIsNone(parents[0].parent_id)
        finally:
            await engine.dispose()


if __name__ == "__main__":
    unittest.main()
