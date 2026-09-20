"""阶段 1 契约：所有格式都产出 Document/Block[]，扫描件不再静默通过。"""

import tempfile
import unittest
from pathlib import Path

from app.collector.pipeline import _FILE_TYPE_MAP, supported_file_types
from app.rag.parsers.factory import ParserFactory, detect_signature


class ParserRegistryTests(unittest.TestCase):
    def test_factory_is_the_single_source_for_the_upload_whitelist(self):
        # 上传白名单、扩展名别名与解析器注册表必须是同一份真相。
        self.assertEqual(ParserFactory.supported_types(), set(_FILE_TYPE_MAP))
        self.assertEqual(supported_file_types(), ParserFactory.supported_types())
        for file_type in sorted(ParserFactory.canonical_types()):
            with self.subTest(file_type=file_type):
                self.assertIsNotNone(ParserFactory.get_parser(file_type))

    def test_aliases_resolve_to_one_canonical_parser(self):
        self.assertEqual(ParserFactory.canonical_type("markdown"), "md")
        self.assertEqual(ParserFactory.canonical_type(".YML"), "txt")
        self.assertEqual(ParserFactory.canonical_type("htm"), "html")
        self.assertIs(ParserFactory.get_parser("md"), ParserFactory.get_parser("markdown"))
        with self.assertRaises(ValueError):
            ParserFactory.get_parser("exe")

    def test_file_signature_detects_mismatched_extension(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fake.docx"
            path.write_text("这不是真正的 docx", encoding="utf-8")
            self.assertEqual(detect_signature(str(path)), "text")
            with self.assertRaises(ValueError) as context:
                ParserFactory.parse_document(
                    str(path), document_id="doc-sig", workspace_id="ws-1", file_type="docx"
                )

        self.assertIn("文件头与扩展名不一致", str(context.exception))


class DocumentParsingTests(unittest.TestCase):
    def _write(self, tmp: str, name: str, text: str) -> str:
        path = Path(tmp) / name
        path.write_text(text, encoding="utf-8")
        return str(path)

    def test_markdown_becomes_structured_blocks(self):
        source = (
            "# RAG 基础\n\n引言段落，说明检索增强生成的整体目标。\n\n"
            "## 检索\n\n- 关键词召回\n- 向量召回\n\n"
            "```python\nprint('hybrid')\n```\n\n"
            "$$E = mc^2$$\n\n"
            "| 组件 | 作用 |\n| --- | --- |\n| Reranker | 重排 |\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = self._write(tmp, "notes.md", source)
            document = ParserFactory.parse_document(
                path, document_id="doc-md", workspace_id="ws-1", file_type="md"
            )

        types = [block.type for block in document.blocks]
        for expected in ("heading", "paragraph", "list", "code", "formula", "table"):
            self.assertIn(expected, types)
        self.assertEqual([block.order for block in document.blocks], list(range(len(document.blocks))))

        heading = next(block for block in document.blocks if block.content == "检索")
        self.assertEqual(heading.heading_level, 2)
        self.assertEqual(heading.section_path, ["RAG 基础", "检索"])
        self.assertEqual(document.title, "RAG 基础")

        formula = next(block for block in document.blocks if block.type == "formula")
        self.assertEqual(formula.latex, "E = mc^2")
        self.assertTrue(formula.is_atomic)

        code = next(block for block in document.blocks if block.type == "code")
        self.assertEqual(code.language, "python")
        self.assertIn("hybrid", code.content)

        table = next(block for block in document.blocks if block.type == "table")
        self.assertEqual(table.table, [["组件", "作用"], ["Reranker", "重排"]])
        # 段落仍保留原始句子，说明结构化没有丢内容
        self.assertTrue(any("引言段落" in block.content for block in document.blocks))
        self.assertGreater(document.text_length(), 0)

    def test_plain_text_becomes_paragraph_blocks(self):
        source = "第一段说明假设。\n\n第二段给出推导。\n"
        with tempfile.TemporaryDirectory() as tmp:
            path = self._write(tmp, "note.txt", source)
            document = ParserFactory.parse_document(
                path, document_id="doc-txt", workspace_id="ws-1", file_type="txt"
            )

        self.assertEqual({block.type for block in document.blocks}, {"paragraph"})
        self.assertEqual(len(document.blocks), 2)
        self.assertFalse(document.quality.scanned)

    def test_csv_becomes_a_native_table_block(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._write(tmp, "params.csv", "组件,作用\nReranker,重排\nChunker,切分\n")
            document = ParserFactory.parse_document(
                path, document_id="doc-csv", workspace_id="ws-1", file_type="csv"
            )

        table = next(block for block in document.blocks if block.type == "table")
        self.assertEqual(table.table[0], ["组件", "作用"])
        self.assertEqual(table.table[2], ["Chunker", "切分"])
        self.assertTrue(table.is_atomic)

    def test_xlsx_keeps_sheet_name_as_heading_and_cells_as_table(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scores.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "参数表"
            sheet.append(["符号", "含义"])
            sheet.append(["P(A|B)", "条件概率"])
            workbook.save(str(path))

            document = ParserFactory.parse_document(
                str(path), document_id="doc-xlsx", workspace_id="ws-1", file_type="xlsx"
            )

        heading = next(block for block in document.blocks if block.type == "heading")
        self.assertEqual(heading.content, "参数表")
        table = next(block for block in document.blocks if block.type == "table")
        # 单元格内的竖线不会再切错列
        self.assertIn(["P(A|B)", "条件概率"], table.table)
        self.assertEqual(table.section_path, ["参数表"])

    def test_html_extracts_headings_paragraphs_lists_code_and_tables(self):
        source = (
            "<html><body><h1>概率论</h1><p>条件概率的定义。</p>"
            "<ul><li>样本空间</li></ul><pre>P(A|B)</pre>"
            "<table><tr><th>符号</th><th>含义</th></tr>"
            "<tr><td>P(A|B)</td><td>条件概率</td></tr></table>"
            "<script>ignore()</script></body></html>"
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = self._write(tmp, "page.html", source)
            document = ParserFactory.parse_document(
                path, document_id="doc-html", workspace_id="ws-1", file_type="html"
            )

        types = {block.type for block in document.blocks}
        self.assertEqual(types, {"heading", "paragraph", "list", "code", "table"})
        table = next(block for block in document.blocks if block.type == "table")
        self.assertEqual(table.table, [["符号", "含义"], ["P(A|B)", "条件概率"]])
        self.assertNotIn("ignore()", " ".join(block.content for block in document.blocks))

    def test_pptx_slide_title_and_body_become_blocks(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "lecture.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "第一节 检索"
            slide.placeholders[1].text = "混合检索把向量与关键词合并。"
            presentation.save(str(path))

            document = ParserFactory.parse_document(
                str(path), document_id="doc-pptx", workspace_id="ws-1", file_type="pptx"
            )

        heading = next(block for block in document.blocks if block.type == "heading")
        self.assertEqual(heading.content, "第一节 检索")
        self.assertEqual(heading.page, 1)
        self.assertTrue(
            any("混合检索" in block.content for block in document.blocks if block.type == "paragraph")
        )

    def test_docx_headings_and_tables_are_typed_blocks(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "textbook.docx"
            docx = DocxDocument()
            docx.add_heading("第一章 概率", level=1)
            docx.add_paragraph("条件概率的定义。")
            table = docx.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "符号"
            table.cell(0, 1).text = "含义"
            table.cell(1, 0).text = "P(A|B)"
            table.cell(1, 1).text = "条件概率"
            docx.save(str(path))

            document = ParserFactory.parse_document(
                str(path), document_id="doc-docx", workspace_id="ws-1", file_type="docx"
            )

        heading = next(block for block in document.blocks if block.type == "heading")
        self.assertEqual(heading.content, "第一章 概率")
        self.assertEqual(heading.heading_level, 1)
        table_block = next(block for block in document.blocks if block.type == "table")
        self.assertIn(["P(A|B)", "条件概率"], table_block.table)
        self.assertEqual(document.quality.signature, "zip")

    def test_scanned_pdf_is_flagged_instead_of_silently_passing(self):
        import fitz

        with tempfile.TemporaryDirectory() as tmp:
            scanned_path = Path(tmp) / "scanned.pdf"
            pdf = fitz.open()
            pdf.new_page()
            pdf.new_page()
            pdf.save(str(scanned_path))
            pdf.close()

            scanned = ParserFactory.parse_document(
                str(scanned_path), document_id="doc-scan", workspace_id="ws-1", file_type="pdf"
            )

            text_path = Path(tmp) / "text.pdf"
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "梯度下降是一种迭代优化方法。")
            pdf.save(str(text_path))
            pdf.close()
            normal = ParserFactory.parse_document(
                str(text_path), document_id="doc-text", workspace_id="ws-1", file_type="pdf"
            )

        self.assertTrue(scanned.quality.scanned)
        self.assertEqual(scanned.quality.degraded, "scanned_pdf")
        self.assertTrue(any("扫描件" in note for note in scanned.quality.notes))
        self.assertEqual(scanned.blocks, [])
        self.assertEqual(scanned.quality.page_count, 2)
        self.assertEqual(scanned.quality.empty_page_count, 2)

        self.assertFalse(normal.quality.scanned)
        self.assertIsNone(normal.quality.degraded)
        self.assertTrue(normal.blocks)
        self.assertGreater(normal.quality.extractable_chars, 0)
        self.assertEqual(normal.quality.signature, "pdf")


if __name__ == "__main__":
    unittest.main()
