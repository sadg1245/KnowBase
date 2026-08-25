"""使用 python-docx 的 DOCX 解析器，用于 KnowBase 采集器流水线。"""

import os
from typing import Any

from loguru import logger

from .base import BaseParser


class DocxParser(BaseParser):
    """解析 .docx 文件，保留标题层级结构和表格内容。

    行为说明
    ---------
    * 段落按标题章节分组为块。每当遇到标题段落时，
      开始一个*新*的块。
    * 连续的正文段落会被拼接（以换行符分隔）到当前块中。
    * 表格会转换为纯文本表示形式并追加到当前块中。
    * 空段落会被静默跳过。
    * ``heading`` 元数据字段始终反映在块内容之前遇到的
      最近的标题文本。
    """

    # --------------------------------------------------------------------- #
    # 辅助方法
    # --------------------------------------------------------------------- #

    @staticmethod
    def _table_to_text(table) -> str:
        """将 python-docx 的 Table 转换为可读的文本块。

        每行转换为管道符分隔的一行::

            | 单元格 A | 单元格 B | 单元格 C |
        """
        lines: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)

    @staticmethod
    def _is_heading(style_name: str | None) -> bool:
        """如果 *style_name* 表示标题样式则返回 True。"""
        if not style_name:
            return False
        return style_name.lower().startswith("heading")

    @staticmethod
    def _heading_level(style_name: str | None) -> int:
        """从样式名称中提取数字标题级别（1-9），无法提取时返回 0。"""
        if not style_name:
            return 0
        lower = style_name.lower()
        if not lower.startswith("heading"):
            return 0
        try:
            return int(lower.replace("heading", "").strip())
        except ValueError:
            return 0

    # --------------------------------------------------------------------- #
    # 公共接口
    # --------------------------------------------------------------------- #

    def parse(self, file_path: str) -> list[dict[str, Any]]:
        """解析 .docx 文件并返回基于章节的文档块。

        参数
        ----------
        file_path : str
            ``.docx`` 文件的路径。

        返回
        -------
        list[dict]
        """
        from docx import Document  # 延迟导入——可选依赖

        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        source_file = os.path.basename(file_path)

        try:
            doc = Document(file_path)
        except Exception as exc:
            logger.error("Failed to open DOCX '{}': {}", source_file, exc)
            raise ValueError(f"Corrupt or unreadable DOCX '{source_file}': {exc}") from exc

        # ------------------------------------------------------------------
        # 按顺序遍历文档正文元素。我们迭代 ``document.element.body``
        # 的子元素，使段落和表格按文档顺序出现
        # （单独使用 ``doc.paragraphs`` 会跳过表格，
        # 而 ``doc.tables`` 会丢失位置上下文）。
        # ------------------------------------------------------------------
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph as DocxParagraph

        chunks: list[dict[str, Any]] = []
        current_heading: str = ""
        current_level: int = 0
        current_lines: list[str] = []
        section_counter: int = 0

        def _flush():
            """将累积的章节保存为一个块。"""
            nonlocal current_lines, section_counter
            text = "\n".join(current_lines).strip()
            if text:
                section_counter += 1
                chunks.append({
                    "content": text,
                    "metadata": {
                        "page_num": section_counter,
                        "heading": current_heading,
                        "source_file": source_file,
                    },
                })
            current_lines = []

        for element in doc.element.body:
            # --- 段落 -------------------------------------------------
            if element.tag.endswith("}p"):
                para = DocxParagraph(element, doc)
                style_name = para.style.name if para.style else None
                text = para.text.strip()

                if not text:
                    continue  # 跳过空段落

                if self._is_heading(style_name):
                    # 在开始新章节前刷新之前的章节
                    _flush()
                    current_heading = text
                    current_level = self._heading_level(style_name)
                    # 标题本身成为新块的第一行
                    current_lines.append(text)
                else:
                    current_lines.append(text)

            # --- 表格 -----------------------------------------------------
            elif element.tag.endswith("}tbl"):
                table = DocxTable(element, doc)
                try:
                    table_text = self._table_to_text(table)
                    if table_text.strip():
                        current_lines.append("")  # 视觉分隔符
                        current_lines.append(table_text)
                        current_lines.append("")
                except Exception as exc:
                    logger.warning(
                        "Could not convert table in '{}': {}", source_file, exc,
                    )

        # 刷新最后一个章节
        _flush()

        if not chunks:
            logger.warning("DOCX '{}' produced no chunks (empty document?).", source_file)

        return chunks
