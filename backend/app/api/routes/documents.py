"""文档上传与管理端点。"""

import os
import uuid
from datetime import datetime, timezone
from typing import Optional

import aiofiles
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from loguru import logger
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user, get_db, get_settings
from app.collector.pipeline import supported_file_types
from app.config import Settings
from app.core.tags import TagValidationError, normalize_tags
from app.models.document import Document
from app.models.chat import DocumentChunk
from app.models.user import User
from app.models.learning import KnowledgePoint
from app.services.document_jobs import enqueue_document_processing, enqueue_learning_generation
from app.schemas.schemas import (
    DocumentSectionDetail,
    DocumentSectionsResponse,
    DocumentResponse,
    DocumentStatusResponse,
    DocumentUpdate,
    UploadResponse,
)
from app.services.ownership import owned_document, owned_workspace


router = APIRouter(tags=["documents"])

# 允许上传的文件扩展名：唯一来源是解析器工厂，避免"上传成功但解析必然失败"
ALLOWED_EXTENSIONS: set[str] = {f".{name}" for name in supported_file_types()}


def _validate_extension(filename: str) -> str:
    """返回小写扩展名，不支持则抛出 HTTPException。"""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"File type '{ext}' is not allowed. "
                f"Accepted types: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            ),
        )
    return ext


async def _validate_workspace(db: AsyncSession, workspace_id: str, user: User):
    """确保知识库存在且属于当前用户。"""
    return await owned_workspace(db, workspace_id, user)


def _normalized_tags(values: list[str] | None) -> list[str]:
    try:
        return normalize_tags(values)
    except TagValidationError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


async def _extract_text(file_path: str, file_type: str) -> str:
    """根据文件类型提取纯文本内容。"""
    import asyncio

    if file_type in (".txt", ".md", ".rst", ".csv", ".json", ".xml", ".yaml", ".yml"):
        import chardet

        async with aiofiles.open(file_path, "rb") as f:
            raw = await f.read()
        detected = chardet.detect(raw)
        encoding = detected.get("encoding") or "utf-8"
        return raw.decode(encoding, errors="replace")

    if file_type in (".html", ".htm"):
        from bs4 import BeautifulSoup

        async with aiofiles.open(file_path, "r", encoding="utf-8", errors="replace") as f:
            html = await f.read()
        soup = BeautifulSoup(html, "html.parser")
        return soup.get_text(separator="\n")

    if file_type == ".pdf":
        from app.collector.content_filter import filter_learning_content
        from app.collector.parsers.pdf_parser import PDFParser

        def _read_pdf() -> str:
            pages = filter_learning_content(PDFParser().parse(file_path))
            return "\n".join(str(page.get("content") or "") for page in pages)

        return await asyncio.to_thread(_read_pdf)

    if file_type == ".docx":
        import docx

        def _read_docx() -> str:
            doc = docx.Document(file_path)
            return "\n".join(para.text for para in doc.paragraphs)

        return await asyncio.to_thread(_read_docx)

    if file_type == ".pptx":
        from pptx import Presentation

        def _read_pptx() -> str:
            prs = Presentation(file_path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
            return "\n".join(texts)

        return await asyncio.to_thread(_read_pptx)

    if file_type == ".xlsx":
        import pandas as pd

        def _read_xlsx() -> str:
            df = pd.read_excel(file_path, engine="openpyxl")
            return df.to_string(index=False)

        return await asyncio.to_thread(_read_xlsx)

    logger.warning("No text extractor for file type '{}'", file_type)
    return ""


# ---------------------------------------------------------------------------
# 上传文档
# ---------------------------------------------------------------------------


@router.post(
    "/workspaces/{workspace_id}/documents",
    response_model=list[UploadResponse],
    status_code=201,
)
async def upload_documents(
    workspace_id: str,
    files: list[UploadFile] = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    settings: Settings = Depends(get_settings),
) -> list[dict]:
    """上传一个或多个文件到工作区。

    每个文件会经过校验、保存到磁盘并排队处理。
    """
    workspace = await _validate_workspace(db, workspace_id, current_user)

    upload_dir = os.path.join(settings.UPLOAD_DIR, workspace_id)
    os.makedirs(upload_dir, exist_ok=True)

    max_bytes = settings.UPLOAD_MAX_SIZE_MB * 1024 * 1024
    responses: list[dict] = []

    for upload_file in files:
        # 校验文件类型
        ext = _validate_extension(upload_file.filename or "unknown")

        # 生成基于 UUID 的文件名并保存
        unique_name = f"{uuid.uuid4().hex}{ext}"
        save_path = os.path.join(upload_dir, unique_name)
        file_size = 0
        try:
            async with aiofiles.open(save_path, "wb") as f:
                while chunk := await upload_file.read(1024 * 1024):
                    file_size += len(chunk)
                    if file_size > max_bytes:
                        raise HTTPException(
                            status_code=413,
                            detail=f"File '{upload_file.filename}' exceeds maximum size of {settings.UPLOAD_MAX_SIZE_MB} MB",
                        )
                    await f.write(chunk)
            if file_size == 0:
                raise HTTPException(status_code=400, detail=f"File '{upload_file.filename}' is empty")
        except Exception:
            if os.path.exists(save_path):
                os.remove(save_path)
            raise

        # 创建数据库记录
        document = Document(
            workspace_id=workspace_id,
            filename=upload_file.filename or unique_name,
            file_path=save_path,
            file_type=ext,
            file_size=file_size,
            chunk_count=0,
            status="pending",
        )
        db.add(document)
        await db.flush()
        await db.refresh(document)

        document.status = "processing"
        await db.flush()
        await db.commit()

        # Dispatch only after another worker session can see the processing state.
        try:
            enqueue_document_processing(document)
            logger.info("Dispatched document '{}' to Celery worker", document.id)
        except Exception as celery_exc:
            document.status = "failed"
            document.error_message = f"Queue dispatch failed: {celery_exc}"[:1000]
            await db.commit()
            logger.warning("Celery dispatch failed for '{}': {}", document.id, celery_exc)
        await db.refresh(document)

        responses.append(
            {
                "document": DocumentResponse.model_validate(document),
                "status": document.status,
            }
        )

        logger.info(
            "Uploaded '{}' -> '{}' for workspace '{}'",
            upload_file.filename,
            unique_name,
            workspace_id,
        )

    return responses


# ---------------------------------------------------------------------------
# 列出工作区中的文档
# ---------------------------------------------------------------------------


@router.get(
    "/workspaces/{workspace_id}/documents",
    response_model=list[DocumentResponse],
)
async def list_documents(
    workspace_id: str,
    status: Optional[str] = Query(None, description="Filter by status"),
    file_type: Optional[str] = Query(None, description="Filter by file type"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[Document]:
    """列出工作区中的所有文档，可按条件筛选。"""
    await _validate_workspace(db, workspace_id, current_user)

    stmt = (
        select(Document)
        .where(Document.workspace_id == workspace_id)
        .order_by(Document.created_at.desc())
    )
    if status:
        stmt = stmt.where(Document.status == status)
    if file_type:
        stmt = stmt.where(Document.file_type == file_type)

    result = await db.execute(stmt)
    documents = result.scalars().all()
    return list(documents)


# ---------------------------------------------------------------------------
# 获取文档详情
# ---------------------------------------------------------------------------


@router.get("/documents/{document_id}", response_model=DocumentResponse)
async def get_document(
    document_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Document:
    """根据 ID 返回单个文档。"""
    return await owned_document(db, document_id, current_user)


@router.patch("/documents/{document_id}", response_model=DocumentResponse)
async def update_document(
    document_id: str,
    payload: DocumentUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Document:
    document = await owned_document(db, document_id, current_user)
    if payload.filename is not None:
        document.filename = payload.filename.strip()
    if payload.tags is not None:
        document.tags = _normalized_tags(payload.tags)
    document.updated_at = datetime.now(timezone.utc)
    await db.commit()
    await db.refresh(document)
    return document


def _section_item(chunk: DocumentChunk) -> dict:
    return {
        "chunk_id": chunk.id,
        "chunk_index": chunk.chunk_index,
        "page_num": chunk.page_num,
        "heading": chunk.heading,
        "heading_level": chunk.heading_level,
        "section_path": list(chunk.section_path or []),
        "content": chunk.content,
    }


async def _document_chunks(db: AsyncSession, document_id: str, user: User) -> list[DocumentChunk]:
    await owned_document(db, document_id, user)
    rows = await db.execute(
        select(DocumentChunk)
        .where(DocumentChunk.document_id == document_id)
        .order_by(DocumentChunk.chunk_index.asc())
    )
    return list(rows.scalars().all())


@router.get("/documents/{document_id}/sections", response_model=DocumentSectionsResponse)
async def list_document_sections(
    document_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    chunks = await _document_chunks(db, document_id, current_user)
    outline: list[dict] = []
    seen: set[tuple] = set()
    for chunk in chunks:
        key = (tuple(chunk.section_path or []), chunk.heading, chunk.page_num)
        if key not in seen:
            seen.add(key)
            outline.append({
                "section_path": list(chunk.section_path or []),
                "heading": chunk.heading,
                "heading_level": chunk.heading_level,
                "page_num": chunk.page_num,
                "chunk_id": chunk.id,
            })
    return {"document_id": document_id, "outline": outline, "items": [_section_item(row) for row in chunks]}


@router.get("/documents/{document_id}/sections/{chunk_id}", response_model=DocumentSectionDetail)
async def get_document_section(
    document_id: str,
    chunk_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    chunks = await _document_chunks(db, document_id, current_user)
    index = next((i for i, row in enumerate(chunks) if row.id == chunk_id), None)
    if index is None:
        raise HTTPException(status_code=404, detail="Document section not found")
    payload = _section_item(chunks[index])
    payload.update({
        "document_id": document_id,
        "previous_chunk_id": chunks[index - 1].id if index > 0 else None,
        "next_chunk_id": chunks[index + 1].id if index + 1 < len(chunks) else None,
    })
    return payload


@router.post("/documents/{document_id}/reprocess", response_model=DocumentResponse, status_code=202)
async def reprocess_document(
    document_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Document:
    document = await owned_document(db, document_id, current_user)
    if document.status in {"pending", "processing"}:
        raise HTTPException(status_code=409, detail="Document processing is already active")
    document.status = "processing"
    document.error_message = None
    document.learning_status = "not_started"
    document.learning_error_message = None
    await db.commit()
    try:
        enqueue_document_processing(document)
    except Exception as exc:
        document.status = "failed"
        document.error_message = f"Queue dispatch failed: {exc}"[:1000]
        await db.commit()
        raise HTTPException(status_code=503, detail=document.error_message) from exc
    await db.refresh(document)
    return document


@router.post("/documents/{document_id}/regenerate-learning", response_model=DocumentResponse, status_code=202)
async def regenerate_document_learning(
    document_id: str,
    overwrite_tags: bool = Query(False, description="明确要求用 AI 标签覆盖人工标签"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> Document:
    document = await owned_document(db, document_id, current_user)
    if document.status != "ready":
        raise HTTPException(status_code=400, detail="Document parsing must be ready first")
    if document.learning_status in {"queued", "generating"}:
        raise HTTPException(status_code=409, detail="Learning generation is already active")
    document.learning_status = "queued"
    document.learning_error_message = None
    await db.commit()
    try:
        enqueue_learning_generation(document.id, overwrite_tags=overwrite_tags)
    except Exception as exc:
        document.learning_status = "failed"
        document.learning_error_message = f"Learning queue dispatch failed: {exc}"[:1000]
        await db.commit()
        raise HTTPException(status_code=503, detail=document.learning_error_message) from exc
    await db.refresh(document)
    return document


@router.get("/documents/{document_id}/content")
async def preview_document(
    document_id: str,
    download: bool = Query(False),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Serve an owned document through the API instead of a public static folder."""
    document = await owned_document(db, document_id, current_user)
    if not os.path.isfile(document.file_path):
        raise HTTPException(status_code=404, detail="Document not found")
    safe_name = document.filename.replace('"', "")
    disposition = "attachment" if download else "inline"
    return FileResponse(
        document.file_path,
        filename=safe_name,
        content_disposition_type=disposition,
    )


# ---------------------------------------------------------------------------
# 获取文档处理状态
# ---------------------------------------------------------------------------


@router.get("/documents/{document_id}/status", response_model=DocumentStatusResponse)
async def get_document_status(
    document_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """仅返回文档的处理状态字段。"""
    document = await owned_document(db, document_id, current_user)
    return {
        "id": document.id,
        "status": document.status,
        "chunk_count": document.chunk_count,
        "error_message": document.error_message,
        "learning_status": document.learning_status,
        "learning_error_message": document.learning_error_message,
        "processed_at": document.processed_at,
        "learning_generated_at": document.learning_generated_at,
    }


# ---------------------------------------------------------------------------
# 删除文档
# ---------------------------------------------------------------------------


@router.delete("/documents/{document_id}", status_code=204)
async def delete_document(
    document_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    settings: Settings = Depends(get_settings),
) -> None:
    """删除文档：从磁盘移除文件、删除数据库记录以及清除 ChromaDB 中的向量。"""
    document = await owned_document(db, document_id, current_user)

    # 从磁盘删除文件
    if os.path.exists(document.file_path):
        os.remove(document.file_path)
        logger.info("Removed file '{}'", document.file_path)

    # 从 ChromaDB 中删除向量
    try:
        from app.core.chroma import get_chroma_client

        chroma_client = get_chroma_client()
        collection_name = f"ws_{document.workspace_id}".replace("-", "_")
        collection = chroma_client.get_collection(name=collection_name)
        # 兼容新旧摄取流水线的元数据字段。
        for metadata_key in ("doc_id", "document_id"):
            try:
                collection.delete(where={metadata_key: document.id})
            except Exception:
                pass
        logger.info(
            "Removed vectors for document '{}' from ChromaDB", document_id
        )
    except Exception as exc:
        logger.warning("Failed to remove vectors from ChromaDB: {}", exc)

    # 删除数据库记录
    await db.execute(delete(KnowledgePoint).where(KnowledgePoint.document_id == document_id))
    await db.delete(document)
    await db.flush()
    logger.info("Deleted document '{}'", document_id)
    return None
